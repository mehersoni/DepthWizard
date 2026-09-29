import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_left_border(cell, color_hex="1D4ED8", sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def generate_formulation_docx():
    doc = Document()

    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 23, 42)     # Slate 900
    COLOR_SECONDARY = RGBColor(29, 78, 216)  # Blue 700
    COLOR_MATH = RGBColor(30, 58, 138)       # Blue 900
    COLOR_TEXT = RGBColor(51, 65, 85)        # Slate 700
    COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT

    # --- Title Page Header ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("DEPTHWIZARD")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Complete Mathematical & Physical Formulation Handbook\nComprehensive Theoretical Foundations, Proofs & Geodetic Algorithms")
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Metadata card
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Author:", "Meher Soni"),
        ("Institution:", "Indian Institute of Technology Madras (IIT Madras)"),
        ("Subject:", "Remote Sensing Photogrammetry & Geodetic Computer Vision"),
        ("Classification:", "Exhaustive Mathematical Handbook across All 9 Pipeline Modules")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY

    def add_equation_box(formula_text, label=""):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8FAFC")
        set_cell_left_border(cell, color_hex="2563EB", sz="28")
        set_cell_margins(cell, 90, 90, 140, 140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        if label:
            r_lbl = p.add_run(f"[{label}]\n")
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = COLOR_SECONDARY

        run = p.add_run(formula_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_MATH
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_impl_note(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"Implementation Reference: {text}")
        run.font.italic = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_MUTED

    # ==================== MODULE 1 ====================
    add_heading_1("1. Monocular Depth & Tiled Spatial Blending")
    add_impl_note("depth/tiled_inference.py")
    doc.add_paragraph(
        "For high-resolution satellite orthophotos exceeding single-pass GPU memory, imagery is partitioned into overlapping tiles "
        "of size T × T with stride S = ⌊T · (1 - overlap)⌋."
    )
    add_heading_2("1.1 2D Hann Window Weighting Matrix")
    doc.add_paragraph(
        "To eliminate seam boundaries across tile overlaps, a 2D Hann weighting window W(y, x) is evaluated as the outer tensor product of two 1D windows:"
    )
    add_equation_box(
        "w_y(i) = sin²( π (2i + 1) / (2H) ),      i ∈ [0, H-1]\n"
        "w_x(j) = sin²( π (2j + 1) / (2W) ),      j ∈ [0, W-1]\n"
        "W(i, j) = w_y(i) ⊗ w_x(j) = w_y(i) · w_x(j)",
        label="Equation 1.1: 2D Hann Window Tensor Formulation"
    )

    add_heading_2("1.2 Multi-Tile Weighted Blending")
    doc.add_paragraph(
        "For M overlapping tile predictions D_k(x, y) positioned at coordinate offsets (x_{0,k}, y_{0,k}):"
    )
    add_equation_box(
        "D_accum(x, y) = Σ [ W_k(x, y) · D_k(x, y) ]\n"
        "W_accum(x, y) = Σ [ W_k(x, y) ]\n"
        "D_final(x, y) = D_accum(x, y) / ( W_accum(x, y) + ε )",
        label="Equation 1.2: Normalized Spatial Disparity Accumulation"
    )

    # ==================== MODULE 2 ====================
    add_heading_1("2. Relative Digital Surface Model Normalization (rDSM)")
    add_impl_note("calibration/rdsm.py")
    doc.add_paragraph(
        "Relative disparity values D(x, y) are normalized to an invariant scale [0.0, 1.0] using robust percentile bounds "
        "to reject sensor anomalies and high-frequency edge outliers:"
    )
    add_equation_box(
        "rDSM(x, y) = clip( (D(x, y) - D_p_low) / (D_p_high - D_p_low + ε),  0.0,  1.0 )\n\n"
        "where:\n"
        "  • D_p_low  = Percentile(D, 0.5%)\n"
        "  • D_p_high = Percentile(D, 99.5%)\n"
        "  • ε        = 10⁻⁷  (prevents division by zero in perfectly flat terrain)",
        label="Equation 2.1: Robust Percentile Normalization"
    )

    # ==================== MODULE 3 ====================
    add_heading_1("3. Metric Elevation Calibration Engine (Absolute Height)")
    add_impl_note("calibration/metric.py, calibration/gcp_calibration.py")
    
    add_heading_2("3.1 Linear Ordinary Least Squares (OLS)")
    doc.add_paragraph(
        "Given K ≥ 2 Ground Control Points P = {(x_i, y_i, H_i)}_{i=1}^K with sampled predicted disparity D_i = D(x_i, y_i):\n"
        "H(x, y) = a · D(x, y) + b"
    )
    add_equation_box(
        "Cost Function:   min_{a, b} L(a, b) = Σ_{i=1}^K ( a · D_i + b - H_i )²\n\n"
        "Closed-Form Solutions:\n"
        "  a = Cov(D, H) / Var(D) = Σ [ (D_i - D_bar)(H_i - H_bar) ] / Σ [ (D_i - D_bar)² ]\n"
        "  b = H_bar - a · D_bar\n\n"
        "where D_bar = (1/K) Σ D_i  and  H_bar = (1/K) Σ H_i.",
        label="Equation 3.1: Closed-Form OLS Metric Calibration"
    )

    add_heading_2("3.2 Single-Point Anchor Calibration (K = 1)")
    doc.add_paragraph("When only a single GCP (D_1, H_1) is available, vertical scale is fixed to a regional prior a_prior:")
    add_equation_box(
        "a = a_prior\n"
        "b = H_1 - a_prior · D_1",
        label="Equation 3.2: Single-Point Datum Anchor"
    )

    add_heading_2("3.3 Prior-Regularized Least Squares (λ-Ridge)")
    doc.add_paragraph(
        "When sparse GCPs (K < 4) have low variance, an L2 scale penalty regularizes a toward the physical prior a_prior:"
    )
    add_equation_box(
        "min_{a, b} [ Σ_{i=1}^K ( a · D_i + b - H_i )² + λ ( a - a_prior )² ]\n\n"
        "Normal Matrix System:\n"
        "[ Σ D_i² + λ     Σ D_i ] [ a ]   [ Σ D_i H_i + λ a_prior ]\n"
        "[ Σ D_i            K   ] [ b ] = [ Σ H_i                 ]\n\n"
        "Determinant:  Δ = K ( Σ D_i² + λ ) - ( Σ D_i )²\n\n"
        "  a = [ K ( Σ D_i H_i + λ a_prior ) - ( Σ D_i )( Σ H_i ) ] / Δ\n"
        "  b = [ ( Σ D_i² + λ )( Σ H_i ) - ( Σ D_i )( Σ D_i H_i + λ a_prior ) ] / Δ",
        label="Equation 3.3: λ-Ridge Analytical Normal Solver"
    )

    add_heading_2("3.4 Non-Linear Polynomial Calibration")
    doc.add_paragraph(
        "For non-linear sensor responses across mountainous terrains (K ≥ d + 1 GCPs):\n"
        "H(x, y) = Σ_{k=0}^d c_k · D(x, y)^k = c_d D^d + c_{d-1} D^{d-1} + ... + c_1 D + c_0"
    )
    add_equation_box(
        "c = ( V^T V )⁻¹ V^T H\n\n"
        "where V_{i, j} = D_i^{d - j} is the Vandermonde design matrix.",
        label="Equation 3.4: Vandermonde Polynomial Solver"
    )

    add_heading_2("3.5 Monotonic Isotonic Regression Calibration")
    doc.add_paragraph(
        "To enforce strict physical elevation monotonicity (∂H / ∂D ≥ 0) without imposing polynomial curvature:"
    )
    add_equation_box(
        "Optimization:   min_{H_hat} Σ_{i=1}^K ( H_i - H_hat_i )²   s.t.  H_hat_i ≤ H_hat_j  whenever  D_i ≤ D_j\n\n"
        "Piecewise Linear Evaluation:\n"
        "H(x, y) = H_hat_k + [ (H_hat_{k+1} - H_hat_k) / (D_{k+1} - D_k) ] · ( D(x, y) - D_k )\n"
        "for D_k ≤ D(x, y) ≤ D_{k+1}.",
        label="Equation 3.5: Monotonic Isotonic Interpolation"
    )

    add_heading_2("3.6 GCP Residuals & Goodness-of-Fit")
    add_equation_box(
        "residual_i = ( a · D_i + b ) - H_i\n"
        "MAE_GCP    = (1/K) Σ_{i=1}^K | residual_i |\n"
        "RMSE_GCP   = √[ (1/K) Σ_{i=1}^K ( residual_i )² ]",
        label="Equation 3.6: Calibration Residual Metrics"
    )

    # ==================== MODULE 4 ====================
    add_heading_1("4. Coarse DEM & SRTM Terrain Datum Anchoring")
    add_impl_note("calibration/srtm_anchor.py")
    
    add_heading_2("4.1 Terrarium RGB Elevation Decoding")
    doc.add_paragraph("AWS / Mapzen terrain tiles encode elevation in 24-bit RGB rasters:")
    add_equation_box(
        "Elevation (m) = ( R · 256.0 + G + B / 256.0 ) - 32768.0",
        label="Equation 4.1: Terrarium RGB Decoding"
    )

    add_heading_2("4.2 Web Mercator & Slippy Map Projection")
    add_equation_box(
        "x_tile = ⌊ (lon + 180°) / 360° · 2^z ⌋\n"
        "y_tile = ⌊ [ 1 - ln( tan(lat_rad) + sec(lat_rad) ) / π ] · ( 2^z / 2 ) ⌋\n\n"
        "EPSG:3857 Metric Coordinates:\n"
        "x_m = lon · ( 20037508.34 / 180° )\n"
        "y_m = [ ln( tan( (90° + lat) π / 360° ) ) / (π / 180°) ] · ( 20037508.34 / 180° )",
        label="Equation 4.2: Slippy Map & Web Mercator Transformation"
    )

    add_heading_2("4.3 Ground Candidate Extraction & Robust Datum Offset")
    add_equation_box(
        "Terrain Mask:   M_terrain = { (x, y)  |  Percentile(D, 1.0%) ≤ D(x, y) ≤ Percentile(D, 25.0%) }\n\n"
        "Datum Offset Anchoring:\n"
        "b = median( H_DEM(x, y) | (x, y) ∈ M_terrain ) - a · median( D(x, y) | (x, y) ∈ M_terrain )\n\n"
        "H(x, y) = a · D(x, y) + b",
        label="Equation 4.3: Robust 25th Percentile Datum Shift"
    )

    # ==================== MODULE 5 ====================
    add_heading_1("5. Shadow Geometry & Vertical Structure Relief Extraction")
    add_impl_note("shadow_detection.py, shadow/m4_physical_raycast_experiment.py")

    add_heading_2("5.1 Spectral Luminance & Rayleigh Skylight Ratios")
    add_equation_box(
        "Luminance:    I = 0.299 R + 0.587 G + 0.114 B\n"
        "BlueRatio:    BlueRatio = B / ( (R + G) / 2.0 + 10⁻⁵ )\n"
        "DiffBR:       DiffBR = ( B - R ) / ( B + R + 10⁻⁴ )",
        label="Equation 5.1: Spectral Shadow Luminance & Rayleigh Ratios"
    )

    add_heading_2("5.2 Continuous Shadow Confidence Formulation")
    add_equation_box(
        "DarkScore     = clip( ( 1.5 · I_thresh - I ) / ( 1.5 · I_thresh + 10⁻⁵ ),  0.0,  1.0 )\n"
        "BlueScore     = clip( ( BlueRatio - 0.80 ) / 0.50,  0.0,  1.0 )\n"
        "SpectralScore = clip( ( DiffBR + 0.10 ) / 0.40,     0.0,  1.0 )\n\n"
        "S_conf(x, y)  = DarkScore × ( 0.60 · BlueScore + 0.40 · SpectralScore )",
        label="Equation 5.2: Multi-Criteria Shadow Confidence"
    )

    add_heading_2("5.3 Shadow Length & Direction Extraction via PCA")
    doc.add_paragraph("For connected component shadow region Ω_k = {(r_i, c_i)}_{i=1}^N:")
    add_equation_box(
        "μ = (1/N) Σ p_i,      Σ = (1/N) Σ ( p_i - μ )( p_i - μ )^T\n\n"
        "Eigendecomposition:   Σ v = λ v  ==>  λ_max,  v_max\n\n"
        "L_pixel  = 4.0 · √[ max( 1.0, λ_max ) ]\n"
        "φ_shadow = atan2( v_max_r,  v_max_c ) · ( 180° / π )",
        label="Equation 5.3: PCA Shadow Axis Eigendecomposition"
    )

    add_heading_2("5.4 Building Height Trigonometry & Absolute Rooftop Elevation")
    add_equation_box(
        "Physical Shadow Length:   L_physical = L_pixel · GSD\n"
        "Building Height:          h_building = L_physical · tan( θ_elev ) = ( L_pixel · GSD ) · tan( θ_elev )\n\n"
        "Absolute Rooftop:         H_rooftop  = H_ground_datum + h_building\n"
        "                                     = H_ground_datum + ( L_pixel · GSD ) · tan( θ_elev )",
        label="Equation 5.4: Physical Shadow Trigonometry"
    )

    add_heading_2("5.5 Multivariate Uncertainty & Error Propagation")
    doc.add_paragraph("For height function H(L_px, s, θ) = L_px · s · tan θ:")
    add_equation_box(
        "Partial Derivatives:\n"
        "  ∂H / ∂L_px = s · tan θ\n"
        "  ∂H / ∂s    = L_px · tan θ\n"
        "  ∂H / ∂θ    = L_px · s · sec² θ = ( L_px · s ) / cos² θ\n\n"
        "Total Analytical Standard Deviation σ_H:\n"
        "σ_H = √[ ( ∂H/∂L_px · δL_px )² + ( ∂H/∂s · δs )² + ( ∂H/∂θ · δθ_rad )² ]\n\n"
        "σ_H = √[ ( s tan θ · δL_px )² + ( L_px tan θ · δs )² + ( (L_px · s / cos² θ) · (π / 180°) · δθ_deg )² ]",
        label="Equation 5.5: Multi-Variable Analytical Error Propagation"
    )

    # ==================== MODULE 6 ====================
    add_heading_1("6. Guided Filter Edge-Preserving Depth Refinement")
    add_impl_note("shadow/guided_filter.py")

    add_heading_2("6.1 Grayscale Guided Filter")
    add_equation_box(
        "Local Linear Model:   q_i = a_k I_i + b_k,   ∀ i ∈ ω_k\n\n"
        "Cost Function:        E(a_k, b_k) = Σ_{i ∈ ω_k} [ ( a_k I_i + b_k - p_i )² + ε a_k² ]\n\n"
        "Parameters:\n"
        "  a_k = [ (1/|ω|) Σ I_i p_i - μ_I μ_p ] / ( σ_I² + ε )\n"
        "  b_k = μ_p - a_k μ_I\n\n"
        "Filtered Output:      q_i = a_bar_i I_i + b_bar_i\n"
        "where a_bar_i = (1/|ω|) Σ_{k ∈ ω_i} a_k   and   b_bar_i = (1/|ω|) Σ_{k ∈ ω_i} b_k.",
        label="Equation 6.1: He et al. Guided Filter Formulation"
    )

    add_heading_2("6.2 3-Channel Color Guided Filter")
    add_equation_box(
        "q_i = a_k^T I_i + b_k = a_{r,k} R_i + a_{g,k} G_i + a_{b,k} B_i + b_k\n\n"
        "a_k = ( Σ_k + ε I_3 )⁻¹ Cov_{I, p}\n"
        "b_k = μ_{p,k} - a_k^T μ_{I,k}\n\n"
        "q_i = a_bar_i^T I_i + b_bar_i",
        label="Equation 6.2: 3-Channel Matrix Covariance Formulation"
    )

    add_heading_2("6.3 Color-Edge Gated Disparity Sharpening")
    add_equation_box(
        "CIELAB Color Gradient Magnitude:\n"
        "|| ∇ I_LAB || = √[ (∂L*/∂x)² + (∂L*/∂y)² + (∂a*/∂x)² + (∂a*/∂y)² + (∂b*/∂x)² + (∂b*/∂y)² ]\n\n"
        "Edge Gating Weight:\n"
        "W_edge = clip( || ∇ I_LAB || / ( Percentile_{95}( || ∇ I_LAB || ) + 10⁻⁵ ),  0.0,  1.0 )\n\n"
        "Sharpened Disparity:\n"
        "D_sharp = D + α_strength · ( D - BoxFilter(D, 3×3) ) · W_edge",
        label="Equation 6.3: Color-Edge Gated Disparity Sharpening"
    )

    # ==================== MODULE 7 ====================
    add_heading_1("7. Geomorphometry & Topographic Shaded Relief")
    add_impl_note("process_image.py")

    add_heading_2("7.1 Surface Gradient & Slope Angle")
    add_equation_box(
        "g_x = ( ∂H / ∂x ) · ( 1 / GSD_x ),      g_y = ( ∂H / ∂y ) · ( 1 / GSD_y )\n\n"
        "Slope (deg) = arctan( √( g_x² + g_y² ) ) · ( 180° / π )    ∈ [0°, 90°]",
        label="Equation 7.1: Surface Gradient & Slope Formulation"
    )

    add_heading_2("7.2 Horn's Algorithm for Shaded Relief (Hillshade)")
    add_equation_box(
        "Terrain Aspect:         Aspect = atan2( -g_y, -g_x )\n"
        "Solar Math Azimuth:     φ_math = radians( 360° - φ_azimuth_geo + 90° )\n\n"
        "Illumination Shading:\n"
        "Shaded    = sin( θ_altitude ) cos( Slope ) + cos( θ_altitude ) sin( Slope ) cos( φ_math - Aspect )\n"
        "Hillshade = ⌊ 255.0 · clip( Shaded, 0.0, 1.0 ) ⌋",
        label="Equation 7.2: Horn's Shaded Relief Formulation"
    )

    add_heading_2("7.3 Heuristic Reliability / Confidence Mapping")
    add_equation_box(
        "Boundary Factor:    C_boundary(x, y) = clip( min( x, W-1-x, y, H-1-y ) / Margin_border,  0.35,  1.0 )\n"
        "Shadow Factor:      C_shadow(x, y)   = 1.0 - 0.35 · S_conf(x, y)\n"
        "Gradient Factor:    C_gradient(x, y) = clip( 1.0 - 0.20 · [ ||∇D|| / ( 2.0 · Percentile_{95}(||∇D||) ) ],  0.50,  1.0 )\n\n"
        "Composite Quality:  C(x, y) = 0.90 · C_boundary(x, y) · C_shadow(x, y) · C_gradient(x, y)    ∈ [0.0, 1.0]",
        label="Equation 7.3: Composite Surface Confidence Formulation"
    )

    # ==================== MODULE 8 ====================
    add_heading_1("8. Statistical Verification & Benchmark Metrics")
    add_impl_note("evaluation/metrics.py")
    doc.add_paragraph(
        "For predicted raster P = {P_i}_{i=1}^N and LiDAR reference DSM R = {R_i}_{i=1}^N:"
    )
    add_equation_box(
        "1. Mean Absolute Error:             MAE  = (1/N) Σ_{i=1}^N | P_i - R_i |\n\n"
        "2. Root Mean Squared Error:         RMSE = √[ (1/N) Σ_{i=1}^N ( P_i - R_i )² ]\n\n"
        "3. Pearson Correlation:             r    = Σ [ (P_i - P_bar)(R_i - R_bar) ] / √[ Σ (P_i - P_bar)² · Σ (R_i - R_bar)² ]\n\n"
        "4. Coefficient of Determination:    R²   = 1.0 - [ Σ ( P_i - R_i )² ] / [ Σ ( R_i - R_bar )² ]\n\n"
        "5. Pixel-wise Absolute Error Map:   E(x, y) = | P(x, y) - R(x, y) |",
        label="Equation 8.1: Complete Benchmark Verification Metrics"
    )

    # ==================== MODULE 9 ====================
    add_heading_1("9. The Zero-Fabrication Safety Contract")
    doc.add_paragraph(
        "To enforce scientific integrity and eliminate hallucinated metric measurements in uncalibrated imagery:"
    )
    add_equation_box(
        "Output Model =\n"
        "  ⎧  H(x, y) = a · D(x, y) + b ∈ ℝ  (metres),       if GCP or DEM calibration supplied\n"
        "  ⎨  H_roof  = H_datum + L · s · tan θ ∈ ℝ (metres),  if solar angles & GSD are present\n"
        "  ⎩  rDSM(x, y) ∈ [0.0, 1.0] (unitless),             if uncalibrated (Zero-Fabrication Mode)",
        label="Equation 9.1: Piecewise Zero-Fabrication Safety Function"
    )

    # Save .docx file
    out_dir = "docs"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "DepthWizard_Mathematical_Formulation_Handbook.docx")
    doc.save(out_path)
    print(f"[Success] Mathematical Formulation Word Document (.docx) saved at: {out_path}")

if __name__ == "__main__":
    generate_formulation_docx()
